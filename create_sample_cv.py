import os
import docx
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

def create_sample_master_cv(filepath: str = "templates/master_cv.docx"):
    """
    Crea un CV Master di esempio con grafica ricca (tabelle a 2 colonne, colori eleganti,
    gerarchia tipografica e bullet point) per testare l'invarianza del layout.
    """
    doc = docx.Document()

    # Margini pagina (1.5 cm)
    for section in doc.sections:
        section.top_margin = Inches(0.6)
        section.bottom_margin = Inches(0.6)
        section.left_margin = Inches(0.6)
        section.right_margin = Inches(0.6)

    # Colori
    NAVY = RGBColor(26, 54, 93)      # #1A365D
    GRAY = RGBColor(100, 116, 139)   # #64748B
    DARK = RGBColor(30, 41, 59)      # #1E293B

    # 1. Header: Nome e Contatti
    p_name = doc.add_paragraph()
    p_name.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run_name = p_name.add_run("MARIO ROSSI")
    run_name.font.name = "Arial"
    run_name.font.size = Pt(22)
    run_name.bold = True
    run_name.font.color.rgb = NAVY

    p_title = doc.add_paragraph()
    run_title = p_title.add_run("Senior Software Engineer & Cloud Architect")
    run_title.font.name = "Arial"
    run_title.font.size = Pt(12)
    run_title.font.color.rgb = GRAY
    run_title.bold = True

    p_contacts = doc.add_paragraph()
    run_c = p_contacts.add_run("Milano, Italia | mario.rossi@email.com | +39 340 1234567 | linkedin.com/in/mariorossi")
    run_c.font.name = "Arial"
    run_c.font.size = Pt(9.5)
    run_c.font.color.rgb = GRAY

    # Linea divisoria
    p_div = doc.add_paragraph()
    run_div = p_div.add_run("—" * 60)
    run_div.font.color.rgb = RGBColor(226, 232, 240)
    run_div.font.size = Pt(8)

    # 2. Tabella a 2 Colonne (Layout tipico e complesso di un CV moderno)
    # Colonna Sinistra (30%): Competenze, Istruzione, Certificazioni
    # Colonna Destra (70%): Profilo Professionale, Esperienze Lavorative
    table = doc.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False

    col_widths = (Inches(2.2), Inches(4.8))
    row = table.rows[0]
    cell_left, cell_right = row.cells[0], row.cells[1]
    cell_left.width = col_widths[0]
    cell_right.width = col_widths[1]

    # --- COLONNA SINISTRA ---
    p_h_skills = cell_left.add_paragraph()
    r = p_h_skills.add_run("COMPETENZE CHIAVE")
    r.font.name = "Arial"
    r.font.size = Pt(11)
    r.bold = True
    r.font.color.rgb = NAVY

    skills = [
        "Python (FastAPI, Django)",
        "Docker & Kubernetes",
        "Cloud (AWS & GCP)",
        "PostgreSQL & Redis",
        "CI/CD Pipelines",
        "Microservices & REST API",
        "Architetture Event-Driven"
    ]
    for s in skills:
        p_s = cell_left.add_paragraph(style='List Bullet')
        p_s.paragraph_format.space_after = Pt(2)
        r_s = p_s.add_run(s)
        r_s.font.name = "Arial"
        r_s.font.size = Pt(9)
        r_s.font.color.rgb = DARK

    # Istruzione
    p_h_edu = cell_left.add_paragraph()
    p_h_edu.paragraph_format.space_before = Pt(12)
    r = p_h_edu.add_run("ISTRUZIONE")
    r.font.name = "Arial"
    r.font.size = Pt(11)
    r.bold = True
    r.font.color.rgb = NAVY

    p_edu = cell_left.add_paragraph()
    r_edu1 = p_edu.add_run("Laurea Magistrale in Informatica\n")
    r_edu1.bold = True
    r_edu1.font.size = Pt(9.5)
    r_edu1.font.color.rgb = DARK
    r_edu2 = p_edu.add_run("Politecnico di Milano (2018 - 2020)")
    r_edu2.font.size = Pt(8.5)
    r_edu2.font.color.rgb = GRAY

    # --- COLONNA DESTRA ---
    # Profilo Professionale
    p_h_sum = cell_right.add_paragraph()
    r = p_h_sum.add_run("PROFILO PROFESSIONALE")
    r.font.name = "Arial"
    r.font.size = Pt(11)
    r.bold = True
    r.font.color.rgb = NAVY

    p_sum = cell_right.add_paragraph()
    p_sum.paragraph_format.space_after = Pt(10)
    r_sum = p_sum.add_run(
        "Software Engineer con oltre 5 anni di esperienza nella progettazione di architetture distribuite "
        "e sistemi backend ad alte prestazioni. Specializzato in soluzioni cloud-native su AWS, automazione DevOps "
        "e rilascio continuo in ambienti Agile collaborativi."
    )
    r_sum.font.name = "Arial"
    r_sum.font.size = Pt(9.5)
    r_sum.font.color.rgb = DARK

    # Esperienza 1
    p_h_exp = cell_right.add_paragraph()
    p_h_exp.paragraph_format.space_before = Pt(8)
    r = p_h_exp.add_run("ESPERIENZA LAVORATIVA")
    r.font.name = "Arial"
    r.font.size = Pt(11)
    r.bold = True
    r.font.color.rgb = NAVY

    p_exp1 = cell_right.add_paragraph()
    r1 = p_exp1.add_run("Tech Solutions S.p.A. | Senior Software Engineer\n")
    r1.bold = True
    r1.font.name = "Arial"
    r1.font.size = Pt(10)
    r1.font.color.rgb = DARK
    r2 = p_exp1.add_run("2021 - Presente | Milano")
    r2.font.name = "Arial"
    r2.font.size = Pt(8.5)
    r2.font.color.rgb = GRAY

    bullets_exp1 = [
        "Guidato la progettazione del backend per la piattaforma di pagamenti, gestendo picchi di 10.000 transazioni al minuto.",
        "Implementato microservizi in Python con FastAPI e code Redis, riducendo la latenza media del 35%.",
        "Configurato cluster Kubernetes su AWS EKS con deployment automatizzato tramite GitHub Actions."
    ]
    for b in bullets_exp1:
        p_b = cell_right.add_paragraph(style='List Bullet')
        p_b.paragraph_format.space_after = Pt(2)
        rb = p_b.add_run(b)
        rb.font.name = "Arial"
        rb.font.size = Pt(9)
        rb.font.color.rgb = DARK

    # Esperienza 2
    p_exp2 = cell_right.add_paragraph()
    p_exp2.paragraph_format.space_before = Pt(8)
    r1 = p_exp2.add_run("Digital Innovation Lab | Software Engineer\n")
    r1.bold = True
    r1.font.name = "Arial"
    r1.font.size = Pt(10)
    r1.font.color.rgb = DARK
    r2 = p_exp2.add_run("2019 - 2021 | Milano")
    r2.font.name = "Arial"
    r2.font.size = Pt(8.5)
    r2.font.color.rgb = GRAY

    bullets_exp2 = [
        "Sviluppato API RESTful in Python/Django integrate con database PostgreSQL.",
        "Migliorato la copertura dei test unitari e di integrazione raggiungendo l'85% di code coverage.",
        "Collaborato a stretto contatto con il team di prodotto in metodologia Scrum bi-settimanale."
    ]
    for b in bullets_exp2:
        p_b = cell_right.add_paragraph(style='List Bullet')
        p_b.paragraph_format.space_after = Pt(2)
        rb = p_b.add_run(b)
        rb.font.name = "Arial"
        rb.font.size = Pt(9)
        rb.font.color.rgb = DARK

    os.makedirs(os.path.dirname(filepath), exist_ok=True)
    doc.save(filepath)
    print(f"✅ Template Master CV generato con successo: {filepath}")

if __name__ == "__main__":
    create_sample_master_cv()
