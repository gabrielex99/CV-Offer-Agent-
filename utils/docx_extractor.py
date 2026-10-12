import docx
from typing import List, Dict, Any

def extract_docx_content(docx_path: str) -> Dict[str, Any]:
    """
    Estrae il testo e la struttura essenziale dal documento DOCX:
    - Lista di paragrafi con eventuale stile
    - Celle e righe delle tabelle
    - Testo complessivo formattato
    """
    doc = docx.Document(docx_path)
    
    paragraphs = []
    for i, p in enumerate(doc.paragraphs):
        text = p.text.strip()
        if text:
            paragraphs.append({
                "index": i,
                "text": text,
                "style": p.style.name if p.style else "Normal"
            })
            
    tables_data = []
    for t_idx, table in enumerate(doc.tables):
        rows_data = []
        for r_idx, row in enumerate(table.rows):
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                rows_data.append(cells)
        if rows_data:
            tables_data.append({
                "table_index": t_idx,
                "rows": rows_data
            })
            
    full_text_lines = [p["text"] for p in paragraphs]
    for t in tables_data:
        for r in t["rows"]:
            full_text_lines.append(" | ".join(r))
            
    return {
        "full_text": "\n".join(full_text_lines),
        "paragraphs": paragraphs,
        "tables": tables_data
    }
